"""Word / PowerPoint / Excel text."""

from __future__ import annotations

import os

from . import Extracted, Segment

MAX_XLSX_ROWS = 20000


def _docx(path: str) -> Extracted:
    import docx

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat
    return Extracted([Segment("\n\n".join(parts))] if parts else [])


def _pptx(path: str) -> Extracted:
    from pptx import Presentation

    prs = Presentation(path)
    segs = []
    for n, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells if c.text.strip()))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append(slide.notes_slide.notes_text_frame.text)
        if texts:
            segs.append(Segment("\n".join(texts), page=n))
    return Extracted(segs, {"slides": len(prs.slides)})


def _xlsx(path: str) -> Extracted:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    segs, rows_seen = [], 0
    try:
        for ws in wb.worksheets:
            lines = []
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append(" | ".join(vals))
                rows_seen += 1
                if rows_seen >= MAX_XLSX_ROWS:
                    break
            if lines:
                segs.append(Segment(f"[sheet {ws.title}]\n" + "\n".join(lines), line=1))
            if rows_seen >= MAX_XLSX_ROWS:
                break
    finally:
        wb.close()
    return Extracted(segs)


def extract_office(path: str) -> Extracted:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext == ".xlsx":
        return _xlsx(path)
    return Extracted()
