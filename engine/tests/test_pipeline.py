"""End-to-end tests with the fake (non-AI) embedder: indexing, exact-sentence search,
page/line locations, filters, incremental updates, moves and deletes.

Run:  uv run python -m unittest discover -s tests -v
(No model download needed. Model quality is measured separately with scripts/eval_golden.py.)
"""

from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from grain_search.chunker import chunk_segment  # noqa: E402
from grain_search.config import Config, load_config  # noqa: E402
from grain_search.crawler import Excluder  # noqa: E402
from grain_search.embedder import HashEmbedder  # noqa: E402
from grain_search.extractors import Segment  # noqa: E402
from grain_search.indexer import run_index  # noqa: E402
from grain_search.search import Searcher, parse_query  # noqa: E402
from grain_search.store import Store  # noqa: E402

SENTENCE = "the invoice was paid late because of the bank holiday in March"
FILLER = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. " * 40


def make_files(root: Path):
    (root / "notes").mkdir(parents=True)
    (root / "photos").mkdir()
    (root / "node_modules" / "pkg").mkdir(parents=True)
    (root / "notes" / "letter.txt").write_text(f"Dear Sam,\n{FILLER}\n\n{SENTENCE}.\nRegards", encoding="utf-8")
    (root / "notes" / "recipe.md").write_text("# Pancakes\nWhisk two eggs with milk and flour.\n", encoding="utf-8")
    (root / "notes" / "budget_2025.csv").write_text("item,cost\nrent,1200\nfood,300\n", encoding="utf-8")
    (root / "notes" / "cp1252.txt").write_bytes("Café crème brûlée recipe".encode("cp1252"))
    (root / "notes" / ".hidden.txt").write_text("secret hidden words", encoding="utf-8")
    (root / "node_modules" / "pkg" / "index.js").write_text("module.exports = 'zebra'", encoding="utf-8")
    (root / "notes" / "binary.txt").write_bytes(b"\x00\x01\x02binary")

    import docx

    d = docx.Document()
    d.add_paragraph("Quarterly report")
    d.add_paragraph("Our penguin colony grew by twelve percent this winter.")
    d.save(root / "notes" / "report.docx")

    from pptx import Presentation

    prs = Presentation()
    for title in ("Intro", "Roadmap", "Volcano safety drill on Thursday"):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
    prs.save(root / "notes" / "deck.pptx")

    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["Supplier", "Amount"])
    wb.active.append(["Acme telescopes", 4200])
    wb.save(root / "notes" / "suppliers.xlsx")

    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(root / "notes" / "manual.pdf"))
    c.drawString(72, 720, "Page one is about installation.")
    c.showPage()
    c.drawString(72, 720, "Never submerge the toaster in the aquarium.")
    c.showPage()
    c.save()

    from PIL import Image

    Image.new("RGB", (64, 48), (200, 20, 20)).save(root / "photos" / "red_car.png")
    Image.new("RGB", (64, 48), (20, 20, 200)).save(root / "photos" / "dog_on_sofa.jpg")
    Image.new("RGB", (64, 48), (20, 200, 20)).save(root / "photos" / "green_forest.webp")
    (root / "photos" / "broken.jpg").write_bytes(b"not really a jpeg")
    Image.new("RGB", (1, 400), (9, 9, 9)).save(root / "photos" / "spacer_line.png")


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.root = self.tmp / "files"
        make_files(self.root)
        self.cfg = Config(roots=[str(self.root)], data_dir=self.tmp / "data")
        self.cfg.excludes = ["*/node_modules/*"]
        self.cfg.image_min_score = 0.0
        self.store = Store(self.cfg.db_path)
        self.emb = HashEmbedder()
        self.rep = run_index(self.cfg, self.store, self.emb)
        self.searcher = Searcher(self.cfg, self.store, self.emb)

    def tearDown(self):
        self.store.close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def names(self, hits):
        return [os.path.basename(h.path) for h in hits]

    # -- indexing ---------------------------------------------------------------
    def test_index_report(self):
        r = self.rep
        self.assertEqual(r.errors, 1, r.error_samples)  # broken.jpg
        broken = self.store.get_file(str(self.root / "photos" / "broken.jpg"))
        self.assertEqual(broken.status, "error")
        self.assertTrue(broken.error.startswith("image: "), broken.error)
        self.assertNotIn("?", broken.error)
        spacer = self.store.get_file(str(self.root / "photos" / "spacer_line.png"))
        self.assertEqual((spacer.status, spacer.error), ("empty", "too small to index (under 32 px)"))
        self.assertEqual(r.images, 3)
        self.assertEqual(r.empty, 2)  # binary.txt + 1-px spacer image
        st = self.store.stats()
        self.assertGreater(st["text_vectors"], 5)
        self.assertEqual(st["model"], "fake/hash-embedder@256")

    def test_hidden_and_excluded_skipped(self):
        paths = " ".join(self.store.all_files())
        self.assertNotIn(".hidden.txt", paths)
        self.assertNotIn("node_modules", paths)

    # -- search -----------------------------------------------------------------
    def test_exact_sentence_is_top1(self):
        res = self.searcher.search(SENTENCE)
        self.assertEqual(self.names(res.docs)[0], "letter.txt")
        top = res.docs[0]
        self.assertTrue(top.exact)
        self.assertIn("bank holiday", top.snippet)
        self.assertEqual(top.line, 4)  # Dear Sam / filler / blank / sentence
        self.assertEqual(res.top.path, top.path)

    def test_quoted_phrase(self):
        res = self.searcher.search('"bank holiday"')
        self.assertEqual(self.names(res.docs)[0], "letter.txt")
        self.assertTrue(res.docs[0].exact)

    def test_missing_text_is_never_labelled_as_contained(self):
        # "invoice" is in letter.txt but "zebra" is in no indexed file
        res = self.searcher.search("invoice zebra")
        self.assertFalse(any(d.contains_text for d in res.docs), [(d.path, d.match) for d in res.docs])

    def test_common_word_does_not_count_as_match(self):
        res = self.searcher.search("the xylophone quartet")
        self.assertFalse(any(d.contains_text for d in res.docs))

    def test_exact_mode_only_returns_files_with_the_text(self):
        res = self.searcher.search("=bank holiday")
        self.assertEqual(self.names(res.docs), ["letter.txt"])
        self.assertEqual(res.images, [])
        self.assertEqual(self.searcher.search("=no such words anywhere").docs, [])

    def test_meaning_mode_skips_text_matching(self):
        res = self.searcher.search("~bank holiday")
        self.assertTrue(all(d.match == "meaning" for d in res.docs))

    def test_related_results_are_capped_when_text_matches_exist(self):
        self.cfg.text_min_score = -1.0  # let everything through the floor
        self.cfg.related_margin = 10.0
        res = self.searcher.search(SENTENCE)
        found = [d for d in res.docs if d.contains_text]
        related = [d for d in res.docs if not d.contains_text]
        self.assertEqual(self.names(found)[0], "letter.txt")
        self.assertLessEqual(len(related), self.cfg.related_limit)
        self.assertTrue(all(d.similarity is not None for d in related))

    def test_pdf_page(self):
        res = self.searcher.search("never submerge the toaster in the aquarium")
        self.assertEqual(self.names(res.docs)[0], "manual.pdf")
        self.assertEqual(res.docs[0].page, 2)

    def test_office_formats(self):
        self.assertEqual(self.names(self.searcher.search("penguin colony grew").docs)[0], "report.docx")
        d = self.searcher.search("volcano safety drill").docs[0]
        self.assertEqual((os.path.basename(d.path), d.page), ("deck.pptx", 3))
        self.assertEqual(self.names(self.searcher.search("Acme telescopes").docs)[0], "suppliers.xlsx")

    def test_cp1252_text(self):
        self.assertEqual(self.names(self.searcher.search("crème brûlée").docs)[0], "cp1252.txt")

    def test_filename_search(self):
        res = self.searcher.search("budget")
        self.assertEqual(self.names(res.docs)[0], "budget_2025.csv")
        self.assertIn("name", res.docs[0].why)
        self.assertEqual(res.docs[0].match, "name")

    def test_image_search(self):
        res = self.searcher.search("red car")
        self.assertEqual(self.names(res.images)[0], "red_car.png")
        self.assertNotIn("red_car.png", self.names(res.docs))

    def test_filters(self):
        self.assertEqual(self.searcher.search("type:img red car").docs, [])
        self.assertEqual(self.searcher.search("type:doc red car").images, [])
        res = self.searcher.search("ext:pdf installation")
        self.assertEqual(set(self.names(res.docs)), {"manual.pdf"})
        res = self.searcher.search(f'in:"{self.root / "photos"}" car')
        self.assertEqual(res.docs, [])
        self.assertEqual(self.names(res.images)[0], "red_car.png")
        self.assertEqual(self.searcher.search("after:2999-01-01 pancakes").docs, [])

    # -- incremental updates -----------------------------------------------------
    def test_rerun_is_noop(self):
        r = run_index(self.cfg, self.store, self.emb)
        self.assertEqual((r.indexed, r.moved, r.removed), (0, 0, 0))
        self.assertEqual(r.unchanged, r.scanned)

    def test_modified_file_reindexed(self):
        p = self.root / "notes" / "recipe.md"
        time.sleep(0.01)
        p.write_text("# Waffles\nPreheat the iron and brush with butter.\n", encoding="utf-8")
        os.utime(p, ns=(time.time_ns(), time.time_ns() + 10**9))
        r = run_index(self.cfg, self.store, self.emb)
        self.assertEqual(r.indexed, 1)
        self.assertEqual(
            self.names(self.searcher.search("preheat the iron and brush with butter").docs)[0], "recipe.md"
        )
        old = [d for d in self.searcher.search('"whisk two eggs"').docs if d.path.endswith("recipe.md")]
        self.assertFalse(old and old[0].exact)

    def test_rename_is_a_move(self):
        old = self.root / "notes" / "letter.txt"
        new = self.root / "notes" / "letter_to_sam.txt"
        old.rename(new)
        r = run_index(self.cfg, self.store, self.emb)
        self.assertEqual((r.moved, r.indexed, r.removed), (1, 0, 0))
        self.assertEqual(self.names(self.searcher.search(SENTENCE).docs)[0], "letter_to_sam.txt")

    def test_delete_removes(self):
        (self.root / "notes" / "report.docx").unlink()
        r = run_index(self.cfg, self.store, self.emb)
        self.assertEqual(r.removed, 1)
        self.assertNotIn("report.docx", self.names(self.searcher.search("penguin colony").docs))

    def test_model_mismatch_refused(self):
        with self.assertRaises(RuntimeError):
            self.store.check_model("google/embeddinggemma-2", 768)


class FailingEmbedder(HashEmbedder):
    def __init__(self, fail_at: str):
        super().__init__()
        self.fail_at = fail_at

    def warm_up(self):
        if self.fail_at == "load":
            raise RuntimeError("Could not load model: CUDA error")

    def encode_documents(self, docs):
        if self.fail_at == "encode":
            raise RuntimeError("out of memory")
        return super().encode_documents(docs)


class ModelFailureTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        make_files(self.tmp / "files")
        self.cfg = Config(roots=[str(self.tmp / "files")], data_dir=self.tmp / "data")
        self.cfg.excludes = [
            "*/node_modules/*"
        ]  # temp dirs live under AppData on Windows, which the default excludes skip

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_load_failure_stops_before_any_file(self):
        store = Store(self.cfg.db_path)
        with self.assertRaisesRegex(RuntimeError, "CUDA error"):
            run_index(self.cfg, store, FailingEmbedder("load"))
        self.assertEqual(store.stats()["files"], 0)
        store.close()

    def test_encode_failure_stops_and_then_recovers(self):
        from grain_search.indexer import ModelError

        store = Store(self.cfg.db_path)
        with self.assertRaisesRegex(ModelError, "out of memory"):
            run_index(self.cfg, store, FailingEmbedder("encode"))
        self.assertEqual(store.stats()["by_status"].get("error", 0), 0)  # files not blamed
        rep = run_index(self.cfg, store, HashEmbedder())  # next run picks everything up
        self.assertGreater(rep.chunks, 5)
        store.close()

    def test_retry_errors(self):
        store = Store(self.cfg.db_path)
        run_index(self.cfg, store, HashEmbedder())
        store.db.execute("UPDATE files SET status='error', error='old model failure' WHERE kind='text'")
        store.commit()
        self.assertEqual(run_index(self.cfg, store, HashEmbedder()).indexed, 0)
        rep = run_index(self.cfg, store, HashEmbedder(), retry_errors=True)
        self.assertGreaterEqual(rep.indexed, 4)
        store.close()


class UnitTests(unittest.TestCase):
    def test_chunk_lines_and_offsets(self):
        text = "\n".join(f"line {i} " + "word " * 30 for i in range(1, 200))
        chunks = chunk_segment(Segment(text, line=1), size=1400, overlap=200)
        self.assertGreater(len(chunks), 5)
        for c in chunks:
            self.assertLessEqual(len(c.text), 1400)
            self.assertTrue(c.text.startswith(f"line {c.line} ") or text.count("\n", 0, c.char_start) + 1 == c.line)
            self.assertEqual(c.text, text[c.char_start : c.char_end].strip())
        # overlap: consecutive chunks share text
        self.assertLess(chunks[1].char_start, chunks[0].char_end)

    def test_parse_query(self):
        q = parse_query('type:img ext:jpg,PNG after:2025-01-01 "red sofa" dog')
        self.assertEqual(q.type, "img")
        self.assertEqual(q.exts, {".jpg", ".png"})
        self.assertIsNotNone(q.after)
        self.assertEqual(q.phrases, ["red sofa"])
        self.assertEqual(q.text, "red sofa dog")
        self.assertEqual(parse_query("note: buy milk").text, "note: buy milk")

    def test_excluder_windows_paths(self):
        ex = Excluder(["C:\\Windows\\*", "*\\node_modules\\*", "*\\AppData\\*"])
        self.assertTrue(ex("C:\\Windows", is_dir=True))
        self.assertTrue(ex("D:\\projects\\app\\node_modules", is_dir=True))
        self.assertTrue(ex("C:\\Users\\me\\AppData\\Local\\x.txt"))
        self.assertFalse(ex("D:\\projects\\app\\src\\main.py"))

    def test_config_yaml(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            (tmp / "c.yaml").write_text(
                "roots: ['~/x']\nmodel:\n  dim: 512\n  modalities: [text]\nchunking:\n  target_tokens: 300\n",
                encoding="utf-8",
            )
            cfg = load_config(tmp / "c.yaml")
            self.assertEqual(cfg.model.dim, 512)
            self.assertFalse(cfg.model.images)
            self.assertEqual(cfg.chunk_chars, 1200)
            self.assertTrue(cfg.roots[0].endswith("x") and "~" not in cfg.roots[0])
        finally:
            shutil.rmtree(tmp)


class CliTest(unittest.TestCase):
    def test_cli_index_search_status(self):
        from grain_search.cli import main

        tmp = Path(tempfile.mkdtemp())
        try:
            make_files(tmp / "files")
            cfg = tmp / "config.yaml"
            cfg.write_text(
                f"roots: ['{(tmp / 'files').as_posix()}']\ndata_dir: '{(tmp / 'data').as_posix()}'\n"
                "excludes: ['*/node_modules/*']\n",
                encoding="utf-8",
            )
            out = io.StringIO()
            with redirect_stdout(out):
                self.assertEqual(main(["--config", str(cfg), "--fake", "index"]), 0)
                self.assertEqual(main(["--config", str(cfg), "--fake", "search", SENTENCE, "-n", "3"]), 0)
                self.assertEqual(main(["--config", str(cfg), "--fake", "status"]), 0)
            text = out.getvalue()
            self.assertIn("letter.txt", text)
            self.assertIn("EXACT", text)
            self.assertTrue((tmp / "data-fake" / "last_results.json").exists())
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
